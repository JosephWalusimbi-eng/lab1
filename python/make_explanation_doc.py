"""Generates docs/Project_and_Code_Explanation.docx: what the project is and how every part of the code works."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
d = Document()
st = d.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)


def H(t, level=1):
    d.add_heading(t, level=level)


def P(t, bold=False, italic=False):
    p = d.add_paragraph()
    r = p.add_run(t)
    r.bold, r.italic = bold, italic
    return p


def B(t, lead=None):
    p = d.add_paragraph(style="List Bullet")
    if lead:
        p.add_run(lead).bold = True
    p.add_run(t)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def CODE(text):
    t = d.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, "F2F2F2")
    c.paragraphs[0].text = ""
    for i, line in enumerate(text.strip("\n").split("\n")):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
    d.add_paragraph()


def TABLE(header, rows, widths=None):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].add_run(h).bold = True
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    d.add_paragraph()


# ------------------------------------------------------------------ title
d.add_heading("WSN Lab 1: Multi-Sensor Edge Node", 0)
P("Project overview and code explanation", bold=True)
P("This document explains what the project is for, how the hardware and software fit together, and what every "
  "part of the code does and why. It is meant to help you understand, present and defend the work.", italic=True)

# ------------------------------------------------------------------ 1
H("1. What the project is about")
P("The lab asks you to build a small, reliable sensor node: an ESP32 microcontroller with several sensors that "
  "collects measurements, checks that they make sense, does some processing on the device itself (\"edge "
  "processing\"), takes a simple local decision, and sends out clean, machine-readable data that Python can "
  "analyse. The scenario is a low-cost node for smart agriculture, environmental monitoring or smart buildings, "
  "and the data becomes the input to later IoT analytics and machine-learning labs.")
P("The lab guide requires you to:")
for x in ["Acquire measurements from at least three sensing modalities (here: temperature/humidity, light, motion).",
          "Detect invalid or implausible values, and never silently delete them.",
          "Calibrate at least one measurement channel (here: temperature, against an LM35 reference).",
          "Filter and extract features locally (5-sample moving average, rate of change, threshold events).",
          "Produce a timestamped, machine-readable data stream (CSV over serial, 1 row per second).",
          "Make a simple local decision or status indicator (NORMAL / MOTION / HIGH_TEMP, LED and buzzer).",
          "Analyse and plot the data in Python (statistics and at least four plots)."]:
    B(x)

H("Why \"edge\" processing?")
P("Instead of sending every raw number to a computer and letting it decide, the ESP32 cleans and interprets the data "
  "itself. This keeps the node useful without a network, reduces the amount of data (and radio energy) needed if "
  "the node later goes wireless, and lets it react instantly, for example by sounding the buzzer.")

# ------------------------------------------------------------------ 2
H("2. System overview")
P("Data flows through four stages:")
TABLE(["Stage", "What happens", "Where"],
      [["1. Sense", "DHT sensor (temperature + humidity), LDR (light), PIR (motion) and an LM35 (reference temperature) "
                    "are read.", "ESP32 pins"],
       ["2. Process", "Validate readings, calibrate, 5-sample moving average, rate of change, threshold with "
                      "hysteresis, event detection.", "ESP32 firmware"],
       ["3. Act / report", "LED and buzzer show the status; one CSV row per second is printed over USB serial.",
        "ESP32 firmware"],
       ["4. Analyse", "Python logs the CSV, calibrates the temperature, computes statistics and makes plots.",
        "PC (Python / Jupyter)"]],
      widths=[1.1, 4.0, 1.4])

# ------------------------------------------------------------------ 3
H("3. Hardware and wiring")
TABLE(["Part", "Role", "ESP32 pin", "Notes"],
      [["DHT temperature/humidity sensor", "Main temperature and humidity measurement", "GPIO4 (digital, single wire)",
        "Slow: maximum about one reading every 2 s. Needs a 10 kOhm pull-up if it is a bare 4-pin part."],
       ["LDR + 1 kOhm divider", "Light level", "GPIO34 (ADC1, analog)",
        "3.3 V - LDR - node - 1 kOhm - GND; the node goes to the pin."],
       ["PIR (HC-SR501)", "Motion detection", "GPIO27 (digital)", "Powered from 5 V (VIN); output is 3.3 V level."],
       ["LM35 (DFRobot Gravity)", "Reference thermometer for calibration", "GPIO35 (ADC1, analog)",
        "Analog output, 10 mV per degC; red = 5 V, black = GND, blue = signal. No library needed."],
       ["LED (+220 ohm)", "Visual status", "GPIO2", "On for any active (non-NORMAL) status."],
       ["Active buzzer", "Audible alert", "GPIO25", "Driven HIGH = sound on."]],
      widths=[1.5, 1.6, 1.5, 2.0])
P("The wiring diagram is in docs/circuit_diagram.png. All sensors share the ESP32 ground. The ESP32's inputs are "
  "3.3 V only, so a 5 V signal must never be connected to a GPIO pin.")
P("Note on the DHT type: the sketch selects the sensor with the line #define DHT_TYPE (currently DHT11). This must "
  "match the physical part: a DHT11 (often blue) and a DHT22 (often white) use the same pin but different "
  "decoding, and a wrong choice gives NaN readings.", italic=True)

# ------------------------------------------------------------------ 4
H("4. The firmware (sensor_node.ino) explained")
P("The sketch has the usual Arduino structure: constants and global variables at the top, helper functions, "
  "setup() which runs once at start-up, and loop() which repeats forever.")

H("4.1 Pins and constants", 2)
CODE("""#define DHT_PIN 4
#define LDR_PIN 34
#define PIR_PIN 27
#define LED_PIN 2
#define BUZZER_PIN 25
#define LM35_PIN 35""")
P("Every pin is named once so the rest of the code reads clearly and a rewire only needs one edit. GPIO34 and "
  "GPIO35 are input-only ADC1 pins, which are also safe to use later when Wi-Fi is on (ADC2 pins do not work "
  "with Wi-Fi).")
P("Tunable settings are grouped as constants:")
TABLE(["Constant", "Value", "Meaning"],
      [["SAMPLE_INTERVAL_MS", "1000", "Main loop runs at 1 Hz (one CSV row per second)."],
       ["DHT_INTERVAL_MS", "2000", "The DHT can only be read about every 2 s."],
       ["PIR_WARMUP_MS", "30000", "PIR is ignored in the status for the first 30 s while it settles."],
       ["TEMP_MIN / TEMP_MAX", "-40 / 80", "Plausible temperature range for validation."],
       ["HUM_MIN / HUM_MAX", "0 / 100", "Plausible humidity range."],
       ["MAX_TEMP_STEP", "5.0 degC", "Largest believable jump between two DHT readings."],
       ["MAX_STEP_REJECTS", "3", "After 3 rejected jumps, accept the new level (the temperature really did change)."],
       ["T_HIGH / T_CLEAR", "35.0 / 34.5", "HIGH_TEMP turns on above 35.0 and off below 34.5 (hysteresis)."],
       ["N_AVG", "5", "Moving-average window length."],
       ["LDR_OVERSAMPLE / LM35_OVERSAMPLE", "8 / 32", "Number of ADC readings averaged per sample."],
       ["CAL_GAIN / CAL_OFFSET", "1.11587 / -1.42758", "Calibration: temp_cal = gain * temp + offset, fitted "
                                             "against the LM35 by calibrate.py (6 points). The Wokwi copy uses "
                                             "1.0 / 0.0."]],
      widths=[2.0, 1.0, 3.5])

H("4.2 Helper functions", 2)
P("plausible(t, h)", bold=True)
CODE("""bool plausible(float t, float h) {
  return !isnan(t) && !isnan(h) &&
         t >= TEMP_MIN && t <= TEMP_MAX &&
         h >= HUM_MIN && h <= HUM_MAX;
}""")
P("Returns true only if both readings are real numbers (the DHT library returns NaN when the read fails) and lie "
  "inside the physically possible range. This is the first layer of data validation.")
P("pushAndAverage(x): the 5-sample moving average", bold=True)
CODE("""float pushAndAverage(float x) {
  avgBuf[avgIdx] = x;
  avgIdx = (avgIdx + 1) % N_AVG;
  if (avgCount < N_AVG) avgCount++;
  float s = 0;
  for (int i = 0; i < avgCount; i++) s += avgBuf[i];
  return s / avgCount;
}""")
P("avgBuf is a circular buffer of the last 5 values. The new value overwrites the oldest one (the index wraps "
  "around with % N_AVG). While fewer than 5 values exist, it averages only what it has (avgCount). Example: for "
  "inputs 30.8, 30.9, 30.8, 30.8, 30.8 the filtered value is 30.82. The average smooths out random noise but reacts "
  "a little late to real changes, about (N-1)/2 samples of lag.")
P("readLM35(): the reference temperature", bold=True)
CODE("""float readLM35() {
  long mvSum = 0;
  for (int i = 0; i < LM35_OVERSAMPLE; i++) mvSum += analogReadMilliVolts(LM35_PIN);
  float t = (mvSum / (float)LM35_OVERSAMPLE) / LM35_MV_PER_C;
  return (t >= 0.0f && t <= 100.0f) ? t : NAN;
}""")
P("The LM35 outputs 10 mV for every degree Celsius, so 25 degC is 250 mV. The function reads the pin 32 times in "
  "millivolts, averages them to cut ADC noise, divides by 10 to get degrees, and returns NaN if the result is "
  "implausible. analogReadMilliVolts() applies the ESP32's factory ADC calibration, and setup() selects the 0 dB "
  "attenuation range (about 0 to 950 mV), which is the most linear range for these small voltages.")
P("printFloat(v, decimals)", bold=True)
P("Prints a number, or the text NaN if it is not a number. This keeps every CSV row full-width even when a "
  "reading is missing.")
P("updateBuzzer(now)", bold=True)
CODE("""void updateBuzzer(uint32_t now) {
  bool on;
  if (buzzerAlarm)  on = (now % (2 * ALARM_HALF_MS)) < ALARM_HALF_MS;
  else if ((int32_t)(beepUntil - now) > 0)  on = true;
  else  on = false;
  digitalWrite(BUZZER_PIN, on ? HIGH : LOW);
}""")
P("This is a non-blocking buzzer: it never waits with delay(). If the alarm flag is set, the buzzer is on for the "
  "first 250 ms of every 500 ms period (now % 500 < 250), giving a steady beep-beep-beep. Otherwise, if a short "
  "chirp was requested (beepUntil is in the future), it stays on until that time passes. It is called at the very "
  "top of loop(), on every pass, so the beeps have their own timing and are not stuck to the 1 Hz sampling.")

H("4.3 setup()", 2)
B("Starts the serial port at 115200 baud.")
B("Configures the pins: PIR and LDR as inputs, LED and buzzer as outputs and switches them off.")
B("Sets the ADC to 12 bits (readings 0 to 4095) and the LM35 pin to 0 dB attenuation.")
B("Starts the DHT library.")
B("Prints comment lines that start with # (ignored by the Python tools) and then the CSV header line.")
B("Waits 2 s for the DHT to settle, then sets nextSample = millis(), the time of the first sample.")

H("4.4 loop(): one pass, step by step", 2)
P("Step 1: timing without drift", bold=True)
CODE("""uint32_t now = millis();
updateBuzzer(now);
if ((int32_t)(now - nextSample) < 0) return;
nextSample += SAMPLE_INTERVAL_MS;""")
P("loop() spins very fast. Everything after the return runs only when the next 1 s slot has arrived. Adding 1000 "
  "to nextSample (instead of using now + 1000) means the time spent inside the loop does not accumulate, so the "
  "rate stays exactly 1 Hz. Casting the difference to a signed int keeps the comparison correct when millis() "
  "overflows after about 49 days.")
P("Step 2: read LDR, LM35 and PIR", bold=True)
P("The LDR is read 8 times and averaged, then divided by 4095 to give a light value between 0 (dark) and 1 "
  "(saturated). With the 1 kOhm divider used in the recorded log it reads close to 1.0 in normal room light. The LM35 goes "
  "through readLM35(). The PIR is a plain digitalRead(): 1 means motion.")
P("Step 3: read the DHT and validate (at most every 2 s)", bold=True)
CODE("""if (!dhtEverRead || (now - lastDhtMs) >= (DHT_INTERVAL_MS - 100)) {
  ...
  float t = dht.readTemperature();
  float h = dht.readHumidity();
  bool ok = plausible(t, h);
  if (ok && haveAccepted && fabsf(t - rawT) > MAX_TEMP_STEP && stepRejects < MAX_STEP_REJECTS) {
    ok = false;  stepRejects++;
  }
  if (ok) { ... calT = CAL_GAIN * rawT + CAL_OFFSET; filtT = pushAndAverage(calT); ... }
  else    { dhtOk = false; invalidCount++; }
}""")
P("Because the loop runs at 1 Hz but the DHT works at about 0.5 Hz, the sensor is read on every second pass. "
  "The rows in between hold the last value and are marked dht_fresh = 0. A reading passes only if it is inside "
  "the plausible range and does not jump more than 5 degC from the last accepted value (a glitch check). After "
  "3 consecutive rejected jumps the new level is accepted, since the temperature may really have moved. When "
  "accepted, the raw value is calibrated (gain and offset), pushed into the moving average, and the rate of "
  "change is computed as the difference from the previous calibrated value. The first sample has no previous "
  "value, so its delta is NaN. A failed read sets dhtOk = false and counts an invalid reading.")
P("Step 4: the local decision", bold=True)
CODE("""if (!dhtOk) status = "INVALID";
else {
  if (!highTemp && calT > T_HIGH) highTemp = true;
  else if (highTemp && calT < T_CLEAR) highTemp = false;
  bool mot = (motion == HIGH) && pirReady;
  if (highTemp && mot) status = "MOTION_AND_HIGH_TEMP";
  else if (highTemp)   status = "HIGH_TEMP";
  else if (mot)        status = "MOTION";
  else                 status = "NORMAL";
  bool active = (strcmp(status, "NORMAL") != 0);
  if (active && !prevActive) event = 1;
  prevActive = active;
}""")
B("Hysteresis: the alarm turns ON above 35.0 degC but only turns OFF below 34.5 degC. Without this band, a "
  "temperature hovering around 35.0 would flip the status on and off every second.", "Threshold with ")
B("The PIR only counts after 30 s (pirReady), because the HC-SR501 gives false triggers while it warms up.",
  "PIR warm-up: ")
B("An event is a single row where the status changes from NORMAL to something else (a rising edge), so a long "
  "period of motion counts as one event, not one per second.", "Event detection: ")
B("If the DHT reading is invalid, the status is INVALID and no decision is made from bad data.", "Invalid data: ")
P("Step 5: LED and buzzer", bold=True)
CODE("""digitalWrite(LED_PIN, (dhtOk && prevActive) ? HIGH : LOW);
buzzerAlarm = dhtOk && highTemp;
if (event && !highTemp) beepUntil = millis() + CHIRP_MS;""")
P("The LED lights for any active status. The buzzer alarm is on while HIGH_TEMP is active and the data is valid. "
  "A single 150 ms chirp is scheduled when a motion-only event happens.")
P("Step 6: print the CSV row", bold=True)
P("A row is always printed with 13 fields, even when data is invalid. Missing values are printed as NaN, never "
  "skipped, so the timeline in Python has no gaps and the lab's rule of never silently deleting measurements is "
  "followed. The timestamp is seconds since boot with 3 decimals.")

H("4.5 Example row", 2)
CODE("22.020,30.90,54.50,1.0000,0,33.05,33.05,0.00,NORMAL,1,1,0,32.49")
P("At 22.02 s after boot: raw temperature 30.90, humidity 54.50 %, light 1.0 (saturated), no motion, calibrated "
  "33.05 (1.11587 * 30.90 - 1.42758), filtered 33.05, change 0.00, status NORMAL, valid, fresh DHT read, not an "
  "event, LM35 reference 32.49 degC. The saved CSV from acquisition.ipynb has one more column, pc_time, "
  "the PC clock time as HH:MM:SS.")

# ------------------------------------------------------------------ 5
H("5. CSV columns")
TABLE(["Column", "Meaning"],
      [["timestamp", "Seconds since boot (3 decimals)."],
       ["temp", "Raw DHT temperature (degC), NaN if invalid."],
       ["humidity", "Relative humidity (%)."],
       ["light", "LDR reading normalised to 0..1."],
       ["motion", "PIR output, 0 or 1 (raw)."],
       ["temp_cal", "Calibrated temperature = CAL_GAIN * temp + CAL_OFFSET."],
       ["temp_filt", "5-sample moving average of temp_cal."],
       ["delta_temp", "temp_cal(t) - temp_cal(t-1) between fresh DHT samples."],
       ["status", "NORMAL, MOTION, HIGH_TEMP, MOTION_AND_HIGH_TEMP or INVALID."],
       ["valid", "1 if the DHT reading passed all checks."],
       ["dht_fresh", "1 if the DHT was read on this row, 0 if the previous value is held."],
       ["event", "1 on the row where the status leaves NORMAL."],
       ["ref_temp", "LM35 reference temperature (degC), used only for calibration."]],
      widths=[1.4, 5.1])
P("The first five columns match the example in the lab guide. Lines starting with # are comments.")

# ------------------------------------------------------------------ 6
H("6. The Python side")
TABLE(["File", "What it does"],
      [["serial_utils.py", "Shared helpers. Opens the serial port, parses each line into a row, and skips comment/"
                           "header lines. Bad lines go to a .log file instead of being lost."],
       ["acquisition.ipynb", "Logs the CSV stream to data/raw/sensor_log_<time>.csv until at least 600 rows are "
                             "collected. Every row is kept, including INVALID ones, and a pc_time column (HH:MM:SS) is added."],
       ["calibrate.py", "collect: averages the DHT and the LM35 over 60 s at several temperatures and saves the "
                        "pairs. fit: least-squares line reference = gain * sensor + offset, reports MAE and RMSE "
                        "before/after (also leave-one-out), saves the model and a figure, and prints CAL_GAIN and "
                        "CAL_OFFSET for the firmware."],
       ["analysis.ipynb", "Loads the newest log, computes the descriptive statistics (count, duration, mean/min/"
                          "max, std, invalid count, events), effective sampling rate and noise reduction of the "
                          "filter, makes the plots (clock-time axis, LM35 overlay, optional PHASES shading), and exports "
                          "processed data."],
       ["make_circuit_diagram.py", "Draws docs/circuit_diagram.png."],
       ["simulation/", "Wokwi project: diagram.json, libraries.txt, and sensor_node_sim/ (the firmware with "
                       "DHT22 and no calibration for the simulator)."],
       ["make_report_draft.py", "Generates the draft technical report docs/lab_report.docx."]],
      widths=[1.8, 4.7])

H("6.1 How calibration works", 2)
P("Calibration compares the DHT's raw reading (x) with the reference (y, the LM35) at several temperatures, then "
  "finds the straight line y = gain * x + offset that fits best. The offset removes a constant bias and the gain "
  "corrects a scale error. MAE (mean absolute error) and RMSE (root mean square error, which penalises large "
  "errors more) are computed before calibration (error of x against y) and after (error of the fitted line "
  "against y). Leave-one-out repeats the fit with each point held out, giving a fairer estimate of the error on "
  "new data. The two numbers are pasted into the firmware and the node then reports temp_cal.")
P("Result: with 6 points between 31.6 and 37.0 degC (one further point was excluded because the temperature was "
  "still changing), the fit is reference = 1.11587 * sensor - 1.42758 with R^2 = 0.957. MAE fell from 2.55 to "
  "0.40 degC (0.60 leave-one-out) and RMSE from 2.60 to 0.46 degC (0.68 leave-one-out). The calibration is valid "
  "only inside that range; outside it the result is an extrapolation.")
P("Limitation: the LM35 itself is accurate to about 0.5 degC and the ESP32 ADC adds error at low voltages, so "
  "the calibration compares the DHT with a second sensor rather than with a traceable thermometer. The overall "
  "accuracy should be quoted as roughly 1 degC.", italic=True)

H("6.2 Plots produced by analysis.ipynb", 2)
for x in ["Temperature vs time (raw, calibrated, and the LM35 reference), with a HH:MM:SS clock-time axis.",
          "Humidity vs time.", "Raw vs 5-sample filtered temperature.", "LDR light vs time.",
          "Motion (PIR) and detected events."]:
    B(x)

H("6.3 Simulation (Wokwi)", 2)
P("Simulation link: https://wokwi.com/projects/476511793297040385. The same firmware runs in Wokwi with three changes (simulation/sensor_node_sim): DHT_TYPE DHT22, because the "
  "simulator's sensor is a DHT22 (with DHT11 decoding it returned wrong values), and gain 1.0 / offset 0.0, "
  "because the fitted calibration belongs to the physical DHT11. An analog NTC sensor stands in for the LM35, so "
  "ref_temp is not a real temperature in the simulator. Wokwi's online compile was delayed at times by a busy "
  "build queue.")

# ------------------------------------------------------------------ 7
H("7. Key concepts in plain words")
TABLE(["Concept", "Explanation"],
      [["Edge processing", "Deciding and cleaning data on the device instead of on a server."],
       ["Oversampling", "Reading an ADC many times and averaging to reduce random noise."],
       ["Moving average", "Average of the last N values; smooths noise but adds lag."],
       ["Rate of change", "How fast a value is changing, x(t) - x(t-1)."],
       ["Hysteresis", "Different on and off thresholds so a signal near the limit does not flicker."],
       ["Event detection", "Counting a change of state once, not every sample while it lasts."],
       ["Validation", "Checking a reading is a number, in range, and not an impossible jump."],
       ["Calibration", "Fitting a correction so readings match a reference."],
       ["MAE / RMSE", "Average size of errors; RMSE weights big errors more."],
       ["Non-blocking timing", "Using millis() comparisons instead of delay() so several things run together."]],
      widths=[1.7, 4.8])

# ------------------------------------------------------------------ 8
H("8. How to run everything end to end")
for i, x in enumerate([
    "Install the ESP32 board package and the libraries DHT sensor library and Adafruit Unified Sensor.",
    "Wire the circuit (section 3), open sensor_node.ino, select the board and port, and upload.",
    "Open the Serial Monitor at 115200 baud and check that CSV rows appear. Close it before using Python.",
    "Calibrate: cd python, run python calibrate.py collect --port COMx --seconds 60 for 5 or more temperatures, "
    "then python calibrate.py fit. Copy CAL_GAIN and CAL_OFFSET into the sketch and upload again.",
    "Record data: run acquisition.ipynb (set PORT) for at least 600 rows, with a 5-minute stable baseline followed "
    "by the LDR and PIR tests. Note the time of each action.",
    "Analyse: run analysis.ipynb. Figures go to figures/ and processed data to data/processed/.",
    "Write the report: fill in docs/lab_report.docx and export it to docs/lab_report.pdf.",
    "Push the repository to GitHub."], 1):
    P(f"{i}. {x}")

# ------------------------------------------------------------------ 9
H("9. Limitations to be aware of")
for x in ["The LDR read 1.0 for the whole recorded log with the 1 kOhm divider resistor, so no light response was "
          "measured; the cause was not established (a wiring fault is likely).",
          "The DHT sensor is slow (about 0.5 Hz) and has limited resolution and accuracy, so the temperature "
          "channel is updated every second row.",
          "The LM35 reference is not a traceable thermometer (about 1 degC overall accuracy), and the DHT11 is only "
          "about +/-2 degC accurate.",
          "The calibration covers only about 31.6-37.0 degC (6 points), so it is not validated at normal room "
          "temperature.",
          "The buzzer is an active type driven with digitalWrite; a passive buzzer would need tone().",
          "The node is wired (USB serial) only: there is no local buffer or MQTT, which are advanced challenges.",
          "The high-temperature threshold is 35 degC for room testing; the lab guide's example uses 30 degC, so set "
          "the value you demonstrate."]:
    B(x)

# ------------------------------------------------------------------ 10
H("10. File map")
CODE("""lab1/
  firmware/sensor_node/sensor_node.ino   ESP32 firmware
  simulation/                            Wokwi project (diagram.json, libraries.txt, wokwi.toml, sensor_node_sim/)
  python/
    serial_utils.py  acquisition.ipynb  calibrate.py  analysis.ipynb
    make_circuit_diagram.py  make_report_draft.py  make_explanation_doc.py
  data/raw  data/processed               raw logs, calibration data, processed outputs
  figures/                               generated plots
  docs/                                  circuit_diagram.png, lab_report.docx/.pdf, this document
  README.md""")

d.save(ROOT / "docs" / "Project_and_Code_Explanation.docx")
print("saved")
